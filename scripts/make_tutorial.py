# -*- coding: utf-8 -*-
"""make_tutorial.py — 生成 docs/教程.xlsx（多页签教程表格，素材来自本仓库 reference/ 与真实案例）。"""
from __future__ import annotations

import os
import re
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REF = os.path.join(ROOT, "reference")
OUT = os.path.join(ROOT, "docs", "教程.xlsx")

H_FILL = PatternFill("solid", fgColor="1F4E79")
H_FONT = Font(bold=True, color="FFFFFF", size=11)
SEC_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def sheet(wb, title, headers, rows, widths=None, freeze="A2", first_col_bold=False):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.fill, cell.font = H_FILL, H_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), ws.max_row)
    if first_col_bold:
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 1).font = Font(bold=True)
    ws.row_dimensions[1].height = 30
    return ws


def read_tsv(name):
    p = os.path.join(REF, name)
    with open(p, encoding="utf-8") as f:
        lines = [l.rstrip("\n").split("\t") for l in f if l.strip()]
    return lines[0], lines[1:]


def parse_copy_samples():
    p = os.path.join(REF, "copy_samples.md")
    blocks, cur = [], None
    for line in open(p, encoding="utf-8"):
        line = line.rstrip("\n")
        m = re.match(r"^## (.+?)  ·  (\S+)  ·  mp=(\S+) lang=(\S+)  ·  数据行 (\d+)", line)
        if m:
            cur = {"file": m.group(1), "lang": m.group(2), "rows": []}
            blocks.append(cur)
            continue
        m = re.match(r"^- \*\*(.+?)\*\*（(\d+) 字符）: (.*)$", line)
        if m and cur is not None:
            cur["rows"].append((cur["lang"], m.group(1), int(m.group(2)), m.group(3)))
    return blocks


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---------------- 00 速查 ----------------
    quick = [
        ("目标", "同一店铺、不同站点、同一 ASIN 链接的商品：把源站点（通常是德国站）上传表里已填的内容，翻译成目标站语言，填进目标站空白上传表，只动数据区，不动表头/格式。"),
        ("输入 1", "源站点上传表（.xlsx/.xlsm，例：惠全优 德国 2399 上传表 .xlsx）"),
        ("输入 2", "目标站点空白模板（例：takini set 模板.xlsm）"),
        ("输入 3", "已有的跨语言取值对照表 reference/vocab_de_eu.tsv（本文档 04 页签）"),
        ("步骤 1", "probe.py 摸清两表结构：站点、语言、属性数、数据行、已填属性 → 确认列是按「属性名」对齐而不是按列号"),
        ("步骤 2", "enums.py / AllowedValues 导出目标模板官方允许值，确认每个枚举字段的合法取值集合"),
        ("步骤 3", "写配置：text（标题/五点/描述/关键词）+ copy（SKU/EAN/价格/型号）+ numeric（价格尺寸）+ vocab（枚举翻译）"),
        ("步骤 4", "fill.py 一次写入。写入方式是 XML 手术：只替换目标 sheet part 与 sharedStrings，其余部件字节级保留（不用 openpyxl 保存，否则丢图片、改 sharedStrings）"),
        ("步骤 5", "check.py 体检：结构 / 必填 / 枚举 / 文案 / 变体关系。FAIL=0 才上传"),
        ("常用命令", "python scripts/probe.py A.xlsx B.xlsm --values --diff"),
        ("", "python scripts/enums.py B.xlsm --pt SWIMWEAR --all --out allowed.tsv"),
        ("", "python scripts/fill.py config/de_to_nl_2399.json"),
        ("", "python scripts/check.py --template B.xlsm --file 产出.xlsm --source A.xlsx --pt SWIMWEAR"),
        ("", "python scripts/vocab.py --pair \"DE|nl_NL|NL\" --pair \"DE|fr_FR|FR\" --out reference/vocab.tsv"),
        ("铁律", "① 枚举值只填目标站 official dropdown 里的值；② 源站点没写的字段，目标站点也不写（尤其合规声明）；③ 父体行不填 size/颜色，子体行必须填 parent_sku。"),
    ]
    sheet(wb, "00-速查", ["项目", "内容"], quick, widths=[14, 150], first_col_bold=True)

    # ---------------- 01 流程 ----------------
    steps = [
        (1, "摸结构", "probe.py 看两表的 marketplace_id / language_tag / 属性数 / 数据行 / 已填属性", "两张表的属性名清单", "两表数据行数一致；确认第 5 行是属性名、第 7 行起是数据"),
        (2, "对列", "按 normalize 后的属性名（删 marketplace_id、language_tag、[]#=）自动建 DE列→目标列 映射", "映射表（脚本自动生成，不手抄列号）", "源表有值、目标表没这列的要列出来单独判断（例：item_shape 荷兰模板没有）"),
        (3, "查枚举", "从 workbook.xml 的 definedNames 找到 SWIMWEAR+属性 指向 Dropdown Lists 的区间，读出官方允许值", "每字段合法取值集合", "脚本必须把 `[ ] # =` 全删掉再拼名字，否则查不到"),
        (4, "翻译取值", "优先用 vocab_de_eu.tsv（历史成交文件反推的实测对照）；vocab 里没有的，人工按官方列表挑，并在配置里写 value_map", "value_map / vocab", "查不到就报错停下，绝不猜一个相近的值"),
        (5, "写文案", "标题≤75、副标题≤125、五点各≤500（第2条≤700）、描述≤2000、关键词≤250；文案结构沿用源站点五段式", "text 段", "主标题品牌开头、无促销词、无全角标点、关键词不重复标题词"),
        (6, "搬硬字段", "SKU / parent_sku / brand / product_id / model / price / dims / weight / fulfillment 原样搬（copy+numeric）", "copy、numeric 段", "价格按数字写；父体行不要有 price/size"),
        (7, "写入", "fill.py 走 XML 手术：替换目标 sheet part + sharedStrings，dimension 扩到末行，其余条目按原 date_time/external_attr 回写", "新 xlsm（与原表同目录，原表不动）", "zip 条目数与空白模板一致（44/44）"),
        (8, "复读校验", "openpyxl 重新打开：表头 1–6 行与空白模板差异必须为 0；源→目标逐格对照缺 0 格多 0 格", "校验输出", "差异 0 才继续"),
        (9, "体检", "check.py：zip/XML/表头/必填/枚举/文案/变体/价格关系", "体检报告", "FAIL=0；WARN 逐条判断（条件必填多为此类目不适用）"),
        (10, "交付", "产出命名 `<目标表名>_<语言>填写_德国<编号>.xlsm`，与目标表同目录；不要用 WPS 二次保存", "可上传文件", "WPS 另存会重排图片部件，文件结构与模板不再一致"),
    ]
    sheet(wb, "01-流程", ["#", "环节", "具体做什么", "产物", "验收标准"], steps, widths=[5, 12, 62, 30, 46])

    # ---------------- 02 站点对照 ----------------
    head, rows = read_tsv("marketplaces.tsv")
    extra = []
    for r in rows:
        mp, lang, sheet_name, country, fname = r
        sizes = {"nl_NL": "S/M/L/XL/XXL 与德文一致", "fr_FR": "Moyen / Grand (l) / TG / TTG (xx_l)",
                 "fr_BE": "同 fr_FR", "es_ES": "同拉丁字母尺码", "en_GB": "S/M/L/XL/XXL 与德文一致",
                 "de_DE": "S (s)/M (m)/L (l)/XL/XXL (xx_l)"}.get(lang, "")
        extra.append((country, lang, mp, sheet_name, fname, sizes))
    sheet(wb, "02-站点对照", ["站点", "language_tag", "marketplace_id", "模板工作表名", "示例模板文件", "尺码写法差异"],
          extra, widths=[10, 14, 18, 14, 34, 34])

    # ---------------- 03 属性处理 ----------------
    proc = []
    with open(os.path.join(REF, "attributes_swimwear_nl.tsv"), encoding="utf-8") as f:
        head = f.readline()
        for line in f:
            a, label, status, col = line.rstrip("\n").split("\t")
            if a in ("item_name#1.value", "title_differentiation#1.value", "product_description#1.value") \
                    or a.startswith(("bullet_point#", "generic_keyword#", "fabric_type#", "model_name#")):
                rule = "text（重写文案）"
            elif a in ("contribution_sku#1.value", "child_parent_sku_relationship#1.parent_sku", "brand#1.value",
                       "amzn1.volt.ca.product_id_value", "model_number#1.value", "manufacturer#1.value",
                       "part_number#1.value", "number_of_pieces#1.value",
                       "dsa_responsible_party_address#1.value",
                       "gpsr_manufacturer_reference#1.gpsr_manufacturer_email_address",
                       "fulfillment_availability#1.fulfillment_channel_code", "product_type#1.value"):
                rule = "copy（原样搬）"
            elif "price" in a or "dimensions" in a or "weight" in a and "battery" not in a:
                rule = "numeric（数字搬）"
            elif status:
                rule = "vocab / value_map（查官方取值）"
            else:
                rule = ""
            if status in ("Vereist", "Erforderlich") and not rule:
                rule = "必填，源表有值就搬"
            proc.append((a, label, status, col, rule))
    sheet(wb, "03-属性处理", ["属性 alias", "荷兰语字段名", "必填级别", "荷兰模板列", "处理方式"],
          proc, widths=[58, 34, 20, 12, 26])

    # ---------------- 04 跨语言取值 ----------------
    head, rows = read_tsv("vocab_de_eu.tsv")
    sheet(wb, "04-跨语言取值", head, rows, widths=[56, 40, 30, 30, 30, 30, 30])

    # ---------------- 05 文案规范 ----------------
    head, rows = read_tsv("char_limits.tsv")
    table = [(r[0], r[1], r[2]) for r in rows]
    table += [
        ("禁用符号", "禁止", "! ? $ _ { } ^ ¬ ¦ ~ # < > * | ; \" \\ / 以及全角标点（：，（）等）"),
        ("促销词", "禁止", "gratis / free shipping / bestseller / topkwaliteit / 100% / korting / sale / günstig / gratuit"),
        ("联系方式", "禁止", "邮箱、URL、电话、WhatsApp、客服字样（DSA/GPSR 专用字段除外，那些字段就是填邮箱/URL 的）"),
        ("HTML", "禁止", "描述与五点只能是纯文本，不能出现 <br> <p> &amp; 等"),
        ("标题结构", "建议", "品牌 + 品类 + 目标人群 + 核心卖点 + 设计点 + 材质/功能"),
        ("关键词策略", "建议", "后台关键词只放标题里没有的同义词/场景词，不要重复标题词，空格分隔，全小写"),
    ]
    sheet(wb, "05-文案规范", ["字段/规则", "上限/性质", "说明"], table, widths=[16, 12, 120], first_col_bold=True)

    # ---------------- 06 验收清单 ----------------
    checks = [
        ("包结构", "zip 条目集合与空白模板一致", "check.py 第一行输出 True"),
        ("XML 完整性", "所有 .xml/.rels 能被 lxml 解析", "异常部件为空"),
        ("表头未破坏", "第 1–6 行与空白模板逐格比对", "差异 0 处"),
        ("硬性必填", "模板 Gegevensdefinities 里标 Vereist 的字段", "每一行都非空"),
        ("枚举合法性", "所有受控文本值都在目标站 official dropdown 内", "越界 0（数字列不参与）"),
        ("源→目标对照", "按 SKU + 属性名逐格比", "缺 0 格 / 多 0 格"),
        ("文案超长", "标题/副标题/五点/描述/关键词", "全部在限内"),
        ("文案违规", "禁用符号/促销词/联系方式/全角", "无"),
        ("变体结构", "父体无 size/颜色；子体有 parent_sku+size+颜色；变体主题一致", "无 FAIL"),
        ("价格关系", "目录价 ≥ 售价；价格是数字", "无 FAIL"),
        ("条件必填", "整列为空的 Voorwaardelijk vereist 字段", "逐条判断是否适用（电池/危险品/尺码区间多不适用）"),
        ("合规声明", "compliance_has_lining 等新增合规字段", "源站点没写就不写，不替客户做声明"),
    ]
    sheet(wb, "06-验收清单", ["检查项", "口径", "通过标准"], checks, widths=[16, 62, 52], first_col_bold=True)

    # ---------------- 07 踩坑 ----------------
    pits = [
        ("copy 被 vocab 劫持", "vocab 是从历史表反推的，里面也有 contribution_sku / model_number / product_id_value（值是 SKU/ASIN），映射先于 copy 走就会报「值映射缺项」并整份不写。", "fill.py 顺序固定 text → copy → vocab/value_map"),
        ("父体行不算必填", "size/颜色/价格/配送/parent_sku 只有子体行才有；按全表计必填会天天刷假 FAIL。", "check.py 用 VARIANT_ONLY + parent_rows 只算子体行"),
        ("条件必填被当硬必填", "ES 模板写 Obligatorio condicionalmente，startswith('obligatorio') 会误判，一次刷 70+ 条假 FAIL。", "先判条件必填再判硬必填，状态词表见 SKILL.md"),
        ("模板被预筛过", "在后台按单个 ASIN 下载的模板，browse node 下拉只剩该商品一个节点（荷兰那份只剩 Bikini's > Sets），拿去上 Tankini/连体款就是错类目。", "enums.py --attr recommended_browse_nodes 看节点数，≤5 就重新下通用模板"),
        ("枚举查不到", "definedName 生成时把 `[ ] # =` 删掉了，拼键时也要删；否则 2674 个名字一个都匹配不上。", "flatfile.AllowedValues.get() 统一处理"),
        ("openpyxl 保存毁文件", "openpyxl 保存 xlsm 会丢 xl/media 里的图片、把 sharedStrings 改成内联字符串、动 drawing 引用。", "改用 zip + XML 手术，只替换 1 个 sheet part 和 sharedStrings"),
        ("WPS 二次保存", "用户在 WPS 里打开再存：条目 44→55，13 张图被重编码改名，还加 docProps/custom.xml。", "交付后不要让 WPS 保存；发现条目数变了就重跑 fill.py"),
        ("文件被占用", "目标文件开着时 os.replace 报 WinError 5 / PermissionError。", "先查 `~$文件名` 锁文件与 EXCEL/WPS 进程，关闭后再写"),
        ("父体行不能有变体值", "父体填 size/color、或子体缺 parent_sku，都会报变体关系错误。", "fill.py 按源表空值跳过，check.py 有专门校验"),
        ("价格被写成文本", "源表价格是文本 '39.99'，直接搬会变成字符串，影响价格解析。", "配置 numeric 列表强制转 float"),
        ("GTIN 豁免", "父体常写 GTIN-vrijgesteld / GTIN Exempt / GTIN-Freistellung，前提是品牌已获批豁免。", "与源站点保持一致，不要自造 GTIN"),
        ("浏览节点每站不同", "同一个商品各站 browse node 名称与 ID 都不一样，不能照抄。", "用 enums.py 取目标站 recommended_browse_nodes 的唯一合法值"),
        ("西班牙站标签不规范", "ES 模板 parentage 取值是 'Principal.' 和 'Niños'（带句号/复数），不要按语义硬翻。", "以 04 页签的实测对照为准"),
        ("后台关键词重复标题", "把标题词再塞进后端，浪费索引位。", "只放同义词/场景词，check.py 会 WARN"),
        ("五点多写一个", "荷兰/德国站模板只有 5 个 bullet_point 列，第 6 个写不进去。", "严格 5 条"),
        ("合规字段别自造", "compliance_has_lining / compliance_has_inner_brief 等是 EU 新增合规声明，属实质申报。", "源站点没写的，目标站点也不写，让客户确认"),
        ("尺码写法差异", "FR 的 M/L/XL 是 Moyen / Grand (l) / TG / TTG (xx_l)，不是 M/L/XL。", "用 04 页签对照，别用字母原文"),
    ]
    sheet(wb, "07-踩坑", ["坑", "现象/原因", "处理"], pits, widths=[18, 76, 66], first_col_bold=True)

    # ---------------- 08 文案样例 ----------------
    rows = []
    for b in parse_copy_samples():
        for lang, field, n, text in b["rows"]:
            rows.append((b["lang"], b["file"], field, n, text))
    sheet(wb, "08-文案样例", ["language_tag", "来源文件", "字段", "字符数", "内容"], rows,
          widths=[12, 34, 18, 8, 110])

    # ---------------- 09 多语言映射速查（只挑常用枚举） ----------------
    head, rows = read_tsv("vocab_de_eu.tsv")
    keep = ("parentage_level", "variation_theme", "amzn1.volt.ca.product_id_type", "material#",
            "item_type_name", "special_feature", "style", "department", "target_gender",
            "age_range_description", "water_resistance_level", "care_instructions", "pattern",
            "weave_type", "fit_type", "swimwear_form_type", "sleeve", "closure", "outer",
            "condition_type", "color#", "shapewear_size#1.size", "recommended_browse_nodes")
    pick = [r for r in rows if any(k in r[0] for k in keep)]
    sheet(wb, "09-枚举速查", head, pick, widths=[56, 42, 30, 30, 30, 30, 30])

    # ---------------- 10 站点专属取值 ----------------
    head, rows = read_tsv("marketplace_specific_values.tsv")
    sheet(wb, "10-站点专属取值", head, rows, widths=[34, 10, 78, 56])

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("[out]", OUT, "sheets:", wb.sheetnames)


if __name__ == "__main__":
    main()
