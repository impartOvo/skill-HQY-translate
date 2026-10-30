# -*- coding: utf-8 -*-
"""fill.py — 按配置把源站点上传表的已填内容翻译/搬运到目标站点上传表（只写数据区，不动表头与格式）。

用法：
    python fill.py config/de_to_nl_2399.json
    python fill.py config.json --output "D:/xx/新文件.xlsm" --dry-run

配置字段见 config/schema.md；核心思路：
  · text      —— 常量文案（标题/五点/描述/关键词/材质…），按属性名派发，所有数据行同值
  · value_map —— 枚举值映射（德语取值 → 目标语官方取值），逐格查表，查不到就报错不瞎填
  · copy      —— 原样复制（SKU/EAN/ASIN/型号/邮箱/价格…）
  · 源表里有值但配置没覆盖的属性会在结尾列出来，避免悄悄丢字段
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import flatfile as ff  # noqa: E402
import openpyxl  # noqa: E402

ff.utf8_stdio()


def load_vocab(path, locale):
    """把 vocab.py 产出的 TSV 读成 {alias: {德文值: 目标语值}}"""
    out = {}
    with open(path, encoding="utf-8") as f:
        header = f.readline().rstrip("\n").split("\t")
        if locale not in header:
            raise SystemExit("vocab 里没有语言 %s（可选: %s）" % (locale, header[2:]))
        idx = header.index(locale)
        for line in f:
            parts = line.rstrip("\n").split("\t")
            if len(parts) <= idx:
                continue
            attr, dv, tv = parts[0], parts[1], parts[idx]
            tv = tv.split(" | ")[0].strip()
            if not tv:
                continue
            out.setdefault(attr, {})[dv] = tv
    return out


def resolve_paths(cfg, cfg_path):
    """① 展开 {{占位符}}：{{SKILL}}=本技能根目录，其余取 cfg['vars'] 或环境变量；
    ② 相对路径先按 config 所在目录、再按上一级目录解析。"""
    skill_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    holder_vars = {"SKILL": skill_root}
    holder_vars.update({k: os.path.expandvars(v) for k, v in (cfg.get("vars") or {}).items()})

    def expand(p):
        if not isinstance(p, str):
            return p
        for k, v in holder_vars.items():
            p = p.replace("{{%s}}" % k, v)
        # 剩下的 {{NAME}} 用同名环境变量兜底
        return re.sub(r"\{\{([A-Za-z0-9_]+)\}\}",
                      lambda m: os.environ.get(m.group(1), m.group(0)), p)

    base = os.path.dirname(os.path.abspath(cfg_path))
    cands = [base, os.path.dirname(base), os.getcwd()]
    for holder, key in ((cfg.get("source") or {}, "path"), (cfg.get("target") or {}, "path"),
                        (cfg.get("vocab") or {}, "path")):
        p = expand(holder.get(key))
        if not p:
            continue
        if not os.path.isabs(p):
            for d in cands:
                if os.path.exists(os.path.join(d, p)):
                    p = os.path.join(d, p)
                    break
        holder[key] = p
    return cfg


def build_records(cfg, verbose=True):
    src = cfg["source"]
    tgt = cfg["target"]
    ws_s = ff.open_source(src["path"])
    wb_t = openpyxl.load_workbook(tgt["path"], keep_vba=tgt["path"].lower().endswith(".xlsm"))
    sheet_t = ff.detect_sheet(wb_t)
    ws_t = wb_t[sheet_t]
    cols_s = ff.attr_map(ws_s, cfg.get("source", {}).get("attr_row", ff.ATTR_ROW))
    cols_t = ff.attr_map(ws_t, cfg.get("target", {}).get("attr_row", ff.ATTR_ROW))
    rows = ff.data_rows(ws_s, cols_s, cfg.get("source", {}).get("data_start", ff.DATA_START))

    text = cfg.get("text", {})
    vmap = {}
    voc = cfg.get("vocab")
    if voc:
        vmap.update(load_vocab(voc["path"], voc["locale"]))
    vmap.update(cfg.get("value_map", {}))
    copy = set(cfg.get("copy", []))
    numeric = set(cfg.get("numeric", []))

    unmapped, missing, bad_value, norecord = [], [], [], set()
    records = []
    for r in rows:
        cells = []
        for a, cs in cols_s.items():
            v = ws_s.cell(r, cs).value
            if ff.is_blank(v):
                continue
            ct = cols_t.get(a)
            if ct is None:
                unmapped.append((a, "目标模板无此属性"))
                continue
            if a in text:
                out = text[a]
            elif a in copy:
                out = v                      # 原样搬（SKU/ASIN/型号/价格/邮箱…）优先于取值映射
            elif a in vmap:
                out = vmap[a].get(ff.norm_value(v))
                if out is None:
                    bad_value.append((r, a, v))
                    continue
            else:
                unmapped.append((a, "配置未覆盖"))
                continue
            isnum = isinstance(out, (int, float)) and not isinstance(out, bool)
            if a in numeric and not isnum:
                try:
                    out = float(str(out).replace(",", "."))
                    isnum = True
                except ValueError:
                    bad_value.append((r, a, "numeric 期望数字，实得 %r" % out))
                    continue
            cells.append((ct, out, isnum))
        if cells:
            records.append((r, sorted(cells)))
        else:
            norecord.add(r)

    stats = {
        "source_rows": len(rows),
        "rows_written": len(records),
        "cells": sum(len(c) for _, c in records),
        "unmapped": sorted(set(unmapped)),
        "bad_value": bad_value,
        "rows_without_data": sorted(norecord),
    }
    if verbose:
        print("[source] %s" % src["path"])
        print("[target] %s" % tgt["path"])
        print("[rows] 源数据行 %d，写入 %d 行，%d 格" % (stats["source_rows"], stats["rows_written"], stats["cells"]))
        if stats["rows_without_data"]:
            print("[warn] 无任何数据的行: %s" % stats["rows_without_data"])
        if stats["unmapped"]:
            print("[warn] 源表有值但未映射的属性（%d）:" % len(stats["unmapped"]))
            for a, why in stats["unmapped"]:
                print("        - %-56s %s" % (a, why))
        if bad_value:
            print("[error] 值映射缺项（%d）——必须补 value_map 或确认源数据:" % len(bad_value))
            for r, a, v in bad_value[:30]:
                print("        - row%d %s = %r" % (r, a, v))
    return records, stats, ws_t, sheet_t


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("config")
    ap.add_argument("--output", default=None)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    cfg = ff.load_config(args.config)
    cfg = resolve_paths(cfg, args.config)
    records, stats, ws_t, sheet_t = build_records(cfg)
    if stats["bad_value"]:
        raise SystemExit(2)
    if args.dry_run:
        print("[dry-run] 未写文件")
        return

    out = args.output or cfg.get("output")
    if not out:
        tgt = cfg["target"]["path"]
        base, ext = os.path.splitext(tgt)
        naming = cfg.get("naming") or {}
        if naming:
            # 统一命名：国家-sku-上传表-dsh生成.xlsm（dsh生成 必须在文件名里，用来区分手工填的表）
            pattern = cfg.get("output_pattern") or "{country}-{sku}-上传表-dsh生成{ext}"
            name = pattern.format(country=naming.get("country", ""),
                                  sku=naming.get("sku") or naming.get("sku_short", ""),
                                  ext=ext)
            out = os.path.join(os.path.dirname(tgt), name)
        else:
            out = base + "_%s%s" % (cfg.get("output_suffix", "filled"), ext)

    info = ff.write_rows_xml(
        target_path=cfg["target"]["path"],
        out_path=out,
        sheet_title=sheet_t,
        records=records,
        max_col=ws_t.max_column,
        last_row=max(r for r, _ in records) if records else ff.DATA_START,
    )
    print("[write] %s  ->  %s" % (info, out))
    print("[ok] 原表未改动；仅替换 %s 与 sharedStrings，其余部件字节级保留" % info["part"])


if __name__ == "__main__":
    main()
