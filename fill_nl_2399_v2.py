# -*- coding: utf-8 -*-
"""德国 2399 上传表 -> 荷兰 takini set 模板：荷兰语翻译填入（XML 级写入，不动表头与格式）。"""
import io, sys, os, re, json, zipfile, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import openpyxl

DE = r"C:\Users\Administrator\AppData\Roaming\ziniaobrowserdatas\ziniao browser\2．惠全优EU\德国\惠全优 德国 2399 上传表 .xlsx"
NL = r"C:\Users\Administrator\AppData\Roaming\ziniaobrowserdatas\ziniao browser\2．惠全优EU\荷兰\takini set 模板.xlsm"
OUT = os.path.join(os.path.dirname(NL), "takini set 模板_荷兰填写_德国2399_v2.xlsm")
WD = r"C:\Users\Administrator\Downloads\10.30新建待办事项"

MPMAP = {"A1PA6795UKMFR9": "MP", "A1805IZSGTT6HS": "MP"}
def norm(s):
    if s is None: return ""
    s = str(s)
    for k, v in MPMAP.items(): s = s.replace(k, v)
    s = re.sub(r"\[language_tag=[A-Za-z_]+", "[language_tag=LANG", s)
    return s.strip()

wde = openpyxl.load_workbook(DE, data_only=True)
sde = wde["Vorlage"]
wnl = openpyxl.load_workbook(NL, keep_vba=True)
snl = wnl["Sjabloon"]

# ---------- 荷兰语文案 ----------
TITLE = "TcIFE Tankini Set Dames Buikbedekkende Zwemkleding High Waist Verstelbaar"
TITLE_DIFF = ("Tweedelig badpak met wikkeldesign en plooien, V-hals, verstelbare bandjes, "
              "zwemkleding voor strand, zwembad en vakantie")
DESC = TITLE + " | " + TITLE_DIFF
KEYWORDS = "badpak bikini zwempak push-up zomer strand vakantie"
MODEL_NAME = "Tankini set"
FABRIC = "80% polyamide, 20% elastaan"
BULLETS = [
    "Materiaal: Hoogwaardig materiaal – BODY: 80% polyamide (nylon), 20% elastaan; voering: 90% polyester, "
    "10% elastaan. De lichte, zachte stof is huidvriendelijk, aangenaam, elastisch en droogt snel in het "
    "water – voor hoog comfort en draagcomfort.",
    "Functie en ontwerp: Tankinisset voor dames, bestaande uit een lang tankinitopje en een bikinibroekje. "
    "Het topje met V-hals en wikkeloptiek, kruisende verstelbare bandjes op de rug en uitneembare softpads "
    "met push-up accentueren de buste. Het bikinibroekje met hoge taille en zijplooien camoufleert de buik "
    "en vormt een slank silhouet. Het lange topje reikt tot over de heupen.",
    "Gelegenheid: Ideaal voor strand, zwembad, vakantie, huwelijksreis, zomer, zonnebaden en alle "
    "wateractiviteiten, perfect bij het zwemmen en zonnebaden. Ook een stijlvol cadeau-idee voor vriendin, "
    "familie en meisjes.",
    "Onderhoudsinstructie: Apart met de hand wassen in koud water, niet chemisch reinigen, geen bleekmiddel "
    "gebruiken, niet uitwringen, aan de lijn laten drogen en direct zonlicht vermijden – zo blijven kleur en "
    "vorm behouden.",
    "Maat: Maten en pasvorm: Verschillende maten beschikbaar: S (EU 36), M (EU 38), L (EU 40), XL (EU 42), "
    "XXL (EU 44). Raadpleeg onze maattabel voor de juiste maat en een optimale pasvorm.",
]
BROWSE_NODE = ("Kleding, schoenen en sieraden > Dames > Kleding > Zwemkleding > Tankini's > Sets "
               "(16392413031)")
COLOR = {"Schwarz": "Zwart", "Grün": "Groen", "Blau": "Blauw", "Lila": "Paars"}

# ---------- 自动按属性名建立列映射 ----------
nl_attr = {}
for c in range(1, snl.max_column + 1):
    v = snl.cell(5, c).value
    if v not in (None, ""):
        nl_attr.setdefault(norm(v), c)
de_attr = {}
for c in range(1, sde.max_column + 1):
    v = sde.cell(5, c).value
    if v not in (None, ""):
        de_attr.setdefault(norm(v), c)

# ---------- 逐列翻译器 ----------
T = {
 2: None, 3: None, 4: {"Erstellen oder Ersetzen (Vollständige Aktualisierung)": "Aanmaken of vervangen (volledige update)",
                       "Bearbeiten (Teilaktualisierung)": "Bewerken (gedeeltelijke update)",
                       "Löschen": "Verwijderen"},
 5: {"Eltern": "Ouder", "Kind": "Kind"},
 6: None, 7: {"GRÖSSE/FARBE": "MAAT/KLEUR"},
 8: ("const", TITLE), 9: ("const", TITLE_DIFF), 10: None,
 11: {"ASIN": "ASIN", "GTIN-Freistellung": "GTIN-vrijgesteld"},
 12: None, 13: ("const", BROWSE_NODE), 18: None, 19: ("const", MODEL_NAME), 20: None,
 39: ("const", DESC),
 40: ("const", BULLETS[0]), 41: ("const", BULLETS[1]), 42: ("const", BULLETS[2]),
 43: ("const", BULLETS[3]), 44: ("const", BULLETS[4]),
 45: ("const", KEYWORDS),
 46: {"Push-Up": "Push-up"}, 47: {"Chlorbeständig": "Chloorbestendig"},
 48: {"Einstellbar": "Verstelbaar"}, 49: {"Schnelltrocknend": "Sneldrogend"},
 50: {"Gepolstert": "Gewatteerd"},
 52: {"Klassisch": "Klassiek"}, 53: {"Damen": "dames"}, 54: {"Weiblich": "Vrouwelijk"},
 55: {"Erwachsener": "Volwassene"},
 56: {"DE/NL/SE/PL": "DE/NL/SE/PL"}, 57: {"Alphanumerisch": "Alfabetisch"}, 58: None,
 62: {"Regular": "Regular"}, 63: {"Regular": "Regular"},
 64: {"Polyester-Mischung": "Polyestermix"}, 67: ("const", FABRIC),
 71: {"Tankini-Set": "Tankini set"}, 72: {"Verstellbar": "Verstelbaar"},
 75: {"Nicht wasserfest": "Niet waterbestendig"}, 77: {"Standard": "Standaard"},
 78: COLOR, 79: COLOR, 81: None, 82: None,
 85: {"Slim": "Slim"}, 86: {"Plain": "Onbewerkt"}, 87: {"Nur Handwäsche": "Alleen handwas"},
 88: {"Hoch": "Hoog"}, 115: {"Einfarbig": "Effen"}, 133: {"V-Ausschnitt": "V-hals"},
 140: {"Tankini": "Tankini"}, 141: {"Ärmellos": "Zonder mouwen"}, 143: {"Pull-on": "Pull-on"},
 157: {"Synthetik": "Synthetisch"}, 163: {"Neu": "Nieuw"},
 165: None, 190: {"AMAZON_EU_RAFN": "AMAZON_EU_RAFN"}, 195: None,
 220: {"Standardvorlage Amazon": "Migrated Template"},
 221: None, 222: {"Zentimeter": "Centimeter", "Zoll": "Inch"},
 223: None, 224: {"Zentimeter": "Centimeter", "Zoll": "Inch"},
 225: None, 226: {"Zentimeter": "Centimeter", "Zoll": "Inch"},
 227: None, 228: {"Kilogramm": "Kilogram", "Gramm": "Gram"},
 229: None, 230: {"Ja": "Ja", "Nein": "Nee"}, 231: {"Ja": "Ja", "Nein": "Nee"},
 246: {"Nicht zutreffend": "Niet van toepassing", "GHS": "GHS", "Lagerung": "Opslag",
       "Transport": "Vervoer", "Entsorgung": "Afval", "Andere": "Overige", "Unbekannt": "Onbekend"},
 289: None, 315: {"Ja": "Ja", "Nein": "Nee"}, 316: None,
}

# 建立 DE 列 -> NL 列
M = {}
skipped = []
for dc in T:
    key = norm(sde.cell(5, dc).value)
    nlc = nl_attr.get(key)
    if nlc is None:
        skipped.append((dc, key))
        continue
    M[dc] = nlc

def conv(spec, v):
    if spec is None:
        return v
    if isinstance(spec, tuple) and spec[0] == "const":
        return spec[1]
    return spec.get(str(v).strip(), "MISS")

# ---------- 生成待写记录 ----------
NUMCOLS = {148, 178, 188, 190, 192, 194}
records = []
problems = []
for r in range(7, sde.max_row + 1):
    cells = []
    for dc, (nlc) in sorted(M.items()):
        v = sde.cell(r, dc).value
        if v is None or (isinstance(v, str) and v.strip() == ""):
            continue
        out = conv(T[dc], v)
        if out == "MISS":
            problems.append((r, dc, nlc, v))
            continue
        isnum = isinstance(out, (int, float)) and not isinstance(out, bool)
        cells.append((nlc, out, isnum))
    if cells:
        records.append((r, sorted(cells)))

print("DE rows:", len(records), " rows span:", records[0][0] if records else None, records[-1][0] if records else None)
print("mapped cols:", len(M), " skipped(no NL column):", skipped)
print("unmapped values:", len(problems))
for p in problems[:20]:
    print("   ", p)

# ---------- 定位 Sjabloon sheet part ----------
z = zipfile.ZipFile(NL)
wbxml = z.read("xl/workbook.xml").decode("utf-8")
rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
rid2t = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
sheets = re.findall(r'<sheet name="([^"]+)"[^>]*r:id="([^"]+)"', wbxml)
name2part = {}
for nm, rid in sheets:
    t = rid2t[rid]
    name2part[nm] = "xl/" + t.lstrip("/").replace("xl/", "")
SHEET = name2part["Sjabloon"]
print("Sjabloon part:", SHEET)

# ---------- sharedStrings ----------
sst = z.read("xl/sharedStrings.xml").decode("utf-8")
count0 = int(re.search(r'<sst[^>]*count="(\d+)"', sst).group(1))
uniq0 = int(re.search(r'<sst[^>]*uniqueCount="(\d+)"', sst).group(1))
existing = re.findall(r"<si>(.*?)</si>", sst, re.S)
index = {}
for i, si in enumerate(existing):
    t = re.sub(r"<[^>]+>", "", si)
    index.setdefault(t, i)
from xml.sax.saxutils import escape
new_si = []
def sid(text):
    if text in index:
        return index[text]
    idx = len(existing) + len(new_si)
    index[text] = idx
    new_si.append('<si><t xml:space="preserve">%s</t></si>' % escape(text))
    return idx

def colname(n):
    s = ""
    while n:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s

rows_xml = []
tot = 0
for r, cells in records:
    parts = ['<row r="%d" spans="1:312">' % r]
    for nlc, out, isnum in cells:
        ref = "%s%d" % (colname(nlc), r)
        if isnum and nlc in NUMCOLS:
            parts.append('<c r="%s"><v>%s</v></c>' % (ref, repr(float(out))))
        else:
            parts.append('<c r="%s" t="s"><v>%d</v></c>' % (ref, sid(str(out))))
        tot += 1
    parts.append("</row>")
    rows_xml.append("".join(parts))

sheet = z.read(SHEET).decode("utf-8")
assert "</sheetData>" in sheet
sheet_new = sheet.replace("</sheetData>", "".join(rows_xml) + "</sheetData>", 1)
sheet_new2 = re.sub(r'<dimension ref="A1:KZ6"', '<dimension ref="A1:KZ27"', sheet_new, count=1)
assert sheet_new2 != sheet_new, "dimension not updated"
sheet_new = sheet_new2

sst_new = sst.replace("</sst>", "".join(new_si) + "</sst>")
sst_new = sst_new.replace('<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="%d"' % count0,
                          '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="%d"' % (count0 + tot), 1)
sst_new = sst_new.replace('uniqueCount="%d"' % uniq0, 'uniqueCount="%d"' % (uniq0 + len(new_si)), 1)

repl = {SHEET: sheet_new.encode("utf-8"), "xl/sharedStrings.xml": sst_new.encode("utf-8")}
tmp = OUT + ".tmp"
with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
    for item in z.infolist():
        data = repl.get(item.filename, z.read(item.filename))
        zi = zipfile.ZipInfo(item.filename, date_time=item.date_time)
        zi.compress_type = zipfile.ZIP_DEFLATED
        zi.external_attr = item.external_attr
        zo.writestr(zi, data)
z.close()
os.replace(tmp, OUT)
print("cells:", tot, "new strings:", len(new_si), "entries kept:", len(zipfile.ZipFile(OUT).namelist()))
print("saved:", OUT)

json.dump({str(k): v for k, v in M.items()}, open(WD + r"\_map.json", "w"), indent=1)
